import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_TAB_ALIGNMENT

NAVY = RGBColor(0x1a, 0x3a, 0x5c)
LIGHT_BLUE = RGBColor(0x2a, 0x6b, 0xa1)
GRAY = RGBColor(0x44, 0x44, 0x44)
BLACK = RGBColor(0x1a, 0x1a, 0x1a)

def add_page_border(sectPr, color_hex="1A3A5C"):
    pgBorders = OxmlElement('w:pgBorders')
    pgBorders.set(qn('w:offsetFrom'), 'page')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '24') # 3 pt width
        border.set(qn('w:space'), '24')
        border.set(qn('w:color'), color_hex)
        pgBorders.append(border)
    sectPr.append(pgBorders)

def parse_hex_color(hex_str, default_rgb):
    try:
        hex_str = hex_str.strip().replace('#', '')
        if len(hex_str) == 6:
            r = int(hex_str[0:2], 16)
            g = int(hex_str[2:4], 16)
            b = int(hex_str[4:6], 16)
            return RGBColor(r, g, b)
    except:
        pass
    return default_rgb

def get_style_config(template_input):
    if isinstance(template_input, dict):
        return template_input
        
    fallback = {
        "font_name": "Calibri",
        "primary_color": "1A3A5C",
        "secondary_color": "2A6BA1",
        "layout_type": "executive",
        "has_border": 1,
        "has_bottom_line": 1
    }
    
    if template_input == "modern":
        fallback.update({
            "font_name": "Arial",
            "primary_color": "10B981",
            "secondary_color": "059669",
            "layout_type": "modern",
            "has_border": 0,
            "has_bottom_line": 0
        })
    elif template_input == "classic":
        fallback.update({
            "font_name": "Times New Roman",
            "primary_color": "1A1A1A",
            "secondary_color": "444444",
            "layout_type": "classic",
            "has_border": 0,
            "has_bottom_line": 1
        })
    return fallback

def setup_doc(with_border=None, template="executive"):
    doc = Document()
    style = get_style_config(template)
    font_name = style['font_name']
    
    if with_border is None:
        with_border = bool(style['has_border'])
        
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        if with_border:
            add_page_border(section._sectPr, style['primary_color'])
            
    st = doc.styles['Normal']
    st.font.name = font_name
    st.font.size = Pt(10)
    st.font.color.rgb = BLACK
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.space_before = Pt(0)
    return doc

def add_bottom_border(p, color='1A3A5C', sz='4'):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

def set_paragraph_shading(p, fill_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    pPr.append(shd)

def heading(doc, text, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    layout = style['layout_type']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    p = doc.add_paragraph()
    
    if layout == "classic":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = primary_rgb
        r.font.name = font_name
        if bool(style['has_bottom_line']):
            add_bottom_border(p, color=style['primary_color'], sz='2')
            
    elif layout == "modern":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        set_paragraph_shading(p, 'E8ECEF')
        r = p.add_run(f"  {text.upper()}")
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = primary_rgb
        r.font.name = font_name
        
    else: # Executive
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = primary_rgb
        r.font.name = font_name
        if bool(style['has_bottom_line']):
            add_bottom_border(p, color=style['primary_color'], sz='4')

def role_header(doc, title, company, dates, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    layout = style['layout_type']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    if layout == "classic":
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(10)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(title)
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = BLACK
        r1.font.name = font_name
        
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(company)
        r2.bold = True; r2.font.size = Pt(10); r2.font.color.rgb = primary_rgb; r2.font.name = font_name
        
        if dates:
            r3 = p2.add_run("  /  ")
            r3.font.size = Pt(10); r3.font.name = font_name
            r4 = p2.add_run(dates)
            r4.italic = True; r4.font.size = Pt(10); r4.font.color.rgb = GRAY; r4.font.name = font_name
            
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        
    elif layout == "modern":
        table = doc.add_table(rows=1, cols=2)
        table.columns[0].width = Inches(5.2)
        table.columns[1].width = Inches(2.0)
        
        p1 = table.cell(0, 0).paragraphs[0]
        p1.paragraph_format.space_before = Pt(8)
        r1 = p1.add_run(title)
        r1.bold = True; r1.font.size = Pt(11); r1.font.color.rgb = primary_rgb; r1.font.name = font_name
        p1.add_run(" | " + company).font.name = font_name
        
        p2 = table.cell(0, 1).paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_before = Pt(8)
        r2 = p2.add_run(dates)
        r2.italic = True; r2.font.size = Pt(10); r2.font.color.rgb = GRAY; r2.font.name = font_name
        
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        
    else: # Executive
        p1 = doc.add_paragraph()
        p1.paragraph_format.space_before = Pt(10)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(title)
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = primary_rgb
        r1.font.name = font_name
        
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        r2 = p2.add_run(company)
        r2.bold = True
        r2.font.size = Pt(10)
        r2.font.color.rgb = BLACK
        r2.font.name = font_name
        
        if dates:
            r3 = p2.add_run("  /  ")
            r3.font.size = Pt(10)
            r3.font.name = font_name
            r4 = p2.add_run(dates)
            r4.italic = True
            r4.font.size = Pt(10)
            r4.font.name = font_name

def bullet(doc, text, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run('\u2022 ' + text.strip())
    r.font.size = Pt(10)
    r.font.name = font_name

def generate_resume(data, output_dir, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    doc = setup_doc(with_border=bool(style['has_border']), template=style)
    
    # Header
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_after = Pt(2)
    r = h.add_run(data.get('name', 'Kwesi Asiedu'))
    r.bold = True; r.font.size = Pt(24); r.font.color.rgb = primary_rgb; r.font.name = font_name
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(4)
    r = c.add_run(data.get('contact', ''))
    r.font.size = Pt(10); r.font.name = font_name; r.font.color.rgb = RGBColor(0x33,0x33,0x33)
    if bool(style['has_bottom_line']):
        add_bottom_border(c, color=style['primary_color'], sz='4' if style['layout_type'] == 'executive' else '2')

    # Clearance
    if data.get('clearance'):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(4)
        if style['layout_type'] == "classic": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('Security Clearance: '); r.bold = True; r.font.size = Pt(10); r.font.name = font_name
        r = p.add_run(data['clearance']); r.font.size = Pt(10); r.font.name = font_name

    # Summary
    if data.get('summary'):
        heading(doc, 'Professional Summary', template=template)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(data['summary']); r.font.size = Pt(10); r.font.name = font_name

    # Core Competencies / Skills
    if data.get('skills'):
        heading(doc, 'Technical Expertise', template=template)
        skills_data = data['skills']
        if isinstance(skills_data, dict):
            for category, items in skills_data.items():
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
                r1 = p.add_run(category + ": ")
                r1.bold = True; r1.font.size = Pt(10); r1.font.name = font_name
                r2 = p.add_run(items)
                r2.font.size = Pt(10); r2.font.name = font_name
        else:
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
            r = p.add_run(skills_data); r.font.size = Pt(10); r.font.name = font_name

    # Experience
    if data.get('roles'):
        heading(doc, 'Professional Experience', template=template)
        for r_data in data['roles']:
            role_header(doc, r_data.get('title',''), r_data.get('company',''), r_data.get('dates',''), template=template)
            for b_text in r_data.get('bullets', []):
                bullet(doc, b_text, template=template)

    # Education
    if data.get('education'):
        heading(doc, 'Education', template=template)
        for edu in data['education']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if ' \u2013 ' in edu or ' - ' in edu:
                separator = ' \u2013 ' if ' \u2013 ' in edu else ' - '
                degree, rest = edu.split(separator, 1)
                r1 = p.add_run(degree.strip() + ' \u2013 ')
                r1.bold = True; r1.font.size = Pt(10); r1.font.name = font_name
                r2 = p.add_run(rest.strip())
                r2.font.size = Pt(10); r2.font.name = font_name
            else:
                r = p.add_run(edu)
                r.bold = True; r.font.size = Pt(10); r.font.name = font_name

    # Certifications
    if data.get('certifications'):
        heading(doc, 'Certifications', template=template)
        for cert in data['certifications']:
            bullet(doc, cert, template=template)

    filename = f"Generated_Resume_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath

def generate_cover_letter(data, output_dir, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    cl = Document()
    for section in cl.sections:
        section.top_margin = Inches(0.75); section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85); section.right_margin = Inches(0.85)
    st = cl.styles['Normal']
    st.font.name = font_name; st.font.size = Pt(11); st.font.color.rgb = BLACK
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)

    def cp(text, bold=False, italic=False, size=11, color=None, sa=6, sb=0):
        p = cl.add_paragraph()
        p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(sb)
        r = p.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font_name
        if color: r.font.color.rgb = color

    cp(data.get('name', 'Kwesi Asiedu'), bold=True, size=16, color=primary_rgb, sa=2)
    cp(data.get('contact', ''), size=10, sa=16, color=GRAY)
    cp(datetime.now().strftime('%B %d, %Y'), sa=14)
    
    cp(data.get('recruiter_name', 'Hiring Manager'), bold=True, sa=0)
    if data.get('company_name'):
        cp(data['company_name'], bold=True, sa=0)
    cp(f"Re: {data.get('job_title', 'Open Position')}", sa=14)
    cp(f"Dear {data.get('recruiter_name', 'Hiring Manager')},", sa=14)

    for para in data.get('cover_letter_paragraphs', []):
        cp(para, sa=11)

    cp('Sincerely,', sa=4, sb=6)
    cp(data.get('name', 'Kwesi Asiedu'), bold=True, sa=0, color=primary_rgb)
    
    filename = f"Generated_CoverLetter_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    cl.save(filepath)
    return filepath

def generate_interview_prep(data, output_dir, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    st = doc.styles['Normal']
    st.font.name = font_name
    st.font.size = Pt(11)
    st.font.color.rgb = BLACK
    st.paragraph_format.space_after = Pt(4)

    def cp(text, bold=False, italic=False, size=11, color=None, sa=6, sb=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(sa)
        p.paragraph_format.space_before = Pt(sb)
        r = p.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font_name
        if color: r.font.color.rgb = color

    # Header
    title = f"Interview Prep: {data.get('job_title', 'Target Role')} at {data.get('company_name', 'Target Company')}"
    cp(title, bold=True, size=16, color=primary_rgb, sa=14)
    cp("Generated based on your tailored resume and the job description.", italic=True, size=10, color=GRAY, sa=16)

    # Questions
    prep_data = data.get('interview_prep', [])
    if not prep_data:
        cp("No interview prep data generated.", sa=10)
    else:
        for i, item in enumerate(prep_data, 1):
            cp(f"Q{i}: {item.get('question', '')}", bold=True, size=12, color=primary_rgb, sa=4, sb=12)
            cp(f"Why they're asking:", bold=True, size=10, color=GRAY, sa=2)
            cp(item.get('rationale', ''), size=10, italic=True, sa=6)
            cp(f"Suggested Answer (STAR):", bold=True, size=10, color=BLACK, sa=2)
            cp(item.get('answer', ''), size=11, sa=8)

    filename = f"Generated_InterviewPrep_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath

def generate_thank_you_note(data, output_dir, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    cl = Document()
    for section in cl.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
    
    st = cl.styles['Normal']
    st.font.name = font_name
    st.font.size = Pt(11)
    st.font.color.rgb = BLACK
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)

    def cp(text, bold=False, italic=False, size=11, color=None, sa=6, sb=0):
        p = cl.add_paragraph()
        p.paragraph_format.space_after = Pt(sa)
        p.paragraph_format.space_before = Pt(sb)
        r = p.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font_name
        if color: r.font.color.rgb = color

    cp(data.get('name', 'Kwesi Asiedu'), bold=True, size=16, color=primary_rgb, sa=2)
    cp(data.get('contact', ''), size=10, sa=16, color=GRAY)
    cp(datetime.now().strftime('%B %d, %Y'), sa=14)
    
    cp(f"Dear {data.get('recruiter_name', 'Interview Team')},", sa=14)

    for para in data.get('thank_you_note', []):
        cp(para, sa=11)

    cp('Best regards,', sa=4, sb=6)
    cp(data.get('name', 'Kwesi Asiedu'), bold=True, sa=0, color=NAVY)
    
    filename = f"Generated_ThankYouNote_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    cl.save(filepath)
    return filepath

def generate_90_day_plan(data, output_dir, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    st = doc.styles['Normal']
    st.font.name = font_name
    st.font.size = Pt(11)
    st.font.color.rgb = BLACK
    st.paragraph_format.space_after = Pt(4)

    def cp(text, bold=False, italic=False, size=11, color=None, sa=6, sb=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(sa)
        p.paragraph_format.space_before = Pt(sb)
        r = p.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font_name
        if color: r.font.color.rgb = color

    # Header
    title = f"90-Day Success Plan: {data.get('job_title', 'Target Role')} at {data.get('company_name', 'Target Company')}"
    cp(title, bold=True, size=16, color=primary_rgb, sa=14)
    cp("A strategic roadmap to deliver immediate value and operational excellence.", italic=True, size=10, color=GRAY, sa=16)

    # Plan Phases
    plan_data = data.get('90_day_plan', [])
    if not plan_data:
        cp("No 90-day plan data generated.", sa=10)
    else:
        for phase_info in plan_data:
            cp(phase_info.get('phase', ''), bold=True, size=12, color=primary_rgb, sa=4, sb=12)
            for goal in phase_info.get('goals', []):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.left_indent = Inches(0.25)
                r = p.add_run('\u2022 ' + goal)
                r.font.size = Pt(10)
                r.font.name = font_name

    filename = f"Generated_90_Day_Plan_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath

def generate_federal_resume(data, output_dir, gs_grade, series, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], BLACK)
    
    doc = setup_doc(with_border=bool(style['has_border']), template=style)
    
    # Header
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_after = Pt(2)
    r = h.add_run(data.get('name', 'Applicant Name'))
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = primary_rgb; r.font.name = font_name
    
    # Contact
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(2)
    r = c.add_run(data.get('contact', ''))
    r.font.size = Pt(10); r.font.name = font_name; r.font.color.rgb = BLACK
    
    if data.get('linkedin'):
        l = doc.add_paragraph(); l.alignment = WD_ALIGN_PARAGRAPH.CENTER; l.paragraph_format.space_after = Pt(2)
        r = l.add_run(f"LinkedIn: {data.get('linkedin', '')}")
        r.font.size = Pt(10); r.font.name = font_name; r.font.color.rgb = BLACK
        
    if data.get('clearance'):
        cl = doc.add_paragraph(); cl.alignment = WD_ALIGN_PARAGRAPH.CENTER; cl.paragraph_format.space_after = Pt(4)
        r = cl.add_run(f"Clearance: {data['clearance']}")
        r.font.size = Pt(10); r.font.name = font_name; r.font.color.rgb = BLACK
        
    # Headline
    if data.get('headline'):
        hl = doc.add_paragraph(); hl.alignment = WD_ALIGN_PARAGRAPH.CENTER; hl.paragraph_format.space_before = Pt(6); hl.paragraph_format.space_after = Pt(6)
        r = hl.add_run(data['headline'])
        r.bold = True; r.font.size = Pt(11); r.font.name = font_name; r.font.color.rgb = BLACK
        
    # Summary
    if data.get('summary'):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
        r = p.add_run(data['summary']); r.font.size = Pt(10); r.font.name = font_name

    # Core Competencies
    if data.get('core_competencies'):
        heading(doc, 'CORE COMPETENCIES', template=template)
        for comp in data['core_competencies']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(comp)
            r.font.size = Pt(10); r.font.name = font_name

    # Technical Skills
    if data.get('technical_skills'):
        heading(doc, 'TECHNICAL SKILLS', template=template)
        skills_data = data['technical_skills']
        for category, items in skills_data.items():
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
            r1 = p.add_run(category + ": ")
            r1.bold = True; r1.font.size = Pt(10); r1.font.name = font_name
            r2 = p.add_run(items)
            r2.font.size = Pt(10); r2.font.name = font_name

    # Professional Experience
    if data.get('roles'):
        heading(doc, 'PROFESSIONAL EXPERIENCE', template=template)
        for r_data in data['roles']:
            title_p = doc.add_paragraph(); title_p.paragraph_format.space_before = Pt(6); title_p.paragraph_format.space_after = Pt(0)
            r = title_p.add_run(r_data.get('title',''))
            r.bold = True; r.font.size = Pt(11); r.font.name = font_name
            
            comp_p = doc.add_paragraph(); comp_p.paragraph_format.space_after = Pt(0)
            r = comp_p.add_run(r_data.get('company_and_location', r_data.get('company', '')))
            r.font.size = Pt(10); r.font.name = font_name
            
            dates_p = doc.add_paragraph(); dates_p.paragraph_format.space_after = Pt(0)
            r = dates_p.add_run(r_data.get('dates',''))
            r.font.size = Pt(10); r.font.name = font_name
            
            if r_data.get('hours'):
                hp = doc.add_paragraph(); hp.paragraph_format.space_after = Pt(0)
                r = hp.add_run(str(r_data['hours']))
                r.font.size = Pt(10); r.font.name = font_name
            if r_data.get('supervisor'):
                sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(0)
                r = sp.add_run(f"Supervisor: {str(r_data['supervisor'])}")
                r.font.size = Pt(10); r.font.name = font_name
            if r_data.get('may_contact'):
                mcp = doc.add_paragraph(); mcp.paragraph_format.space_after = Pt(6)
                r = mcp.add_run(f"May Contact Supervisor: {str(r_data['may_contact'])}")
                r.font.size = Pt(10); r.font.name = font_name

            dp = doc.add_paragraph(); dp.paragraph_format.space_after = Pt(2)
            r = dp.add_run("Duties & Accomplishments")
            r.bold = True; r.font.size = Pt(10); r.font.name = font_name

            for b_text in r_data.get('bullets', []):
                bullet(doc, str(b_text), template=template)
                
            if r_data.get('key_technologies'):
                ktp = doc.add_paragraph(); ktp.paragraph_format.space_before = Pt(4); ktp.paragraph_format.space_after = Pt(6)
                r1 = ktp.add_run("Key Technologies: ")
                r1.bold = True; r1.font.size = Pt(10); r1.font.name = font_name
                r2 = ktp.add_run(str(r_data['key_technologies']))
                r2.font.size = Pt(10); r2.font.name = font_name

    # Education
    if data.get('education'):
        heading(doc, 'Education', template=template)
        for edu in data['education']:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if ' \u2013 ' in edu or ' - ' in edu:
                separator = ' \u2013 ' if ' \u2013 ' in edu else ' - '
                degree, rest = edu.split(separator, 1)
                r1 = p.add_run(degree.strip() + ' \u2013 ')
                r1.bold = True; r1.font.size = Pt(10); r1.font.name = font_name
                r2 = p.add_run(rest.strip())
                r2.font.size = Pt(10); r2.font.name = font_name
            else:
                r = p.add_run(edu)
                r.bold = True; r.font.size = Pt(10); r.font.name = font_name

    # Certifications
    if data.get('certifications'):
        heading(doc, 'Certifications', template=template)
        for cert in data['certifications']:
            bullet(doc, cert, template=template)

    filename = f"Federal_Resume_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath

def generate_ksa_narrative(data, output_dir):
    doc = setup_doc()
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(f"KSA Narratives - {data.get('name', '')}")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    add_bottom_border(h)
    
    for ksa in data.get('ksas', []):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        r = p.add_run("KSA Prompt: " + ksa.get('prompt', ''))
        r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
        
        n = doc.add_paragraph()
        n.add_run(ksa.get('narrative', ''))
        
    filename = f"Federal_KSAs_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath

def generate_ats_report(data, output_dir, template="executive"):
    style = get_style_config(template)
    font_name = style['font_name']
    primary_rgb = parse_hex_color(style['primary_color'], NAVY)
    secondary_rgb = parse_hex_color(style['secondary_color'], LIGHT_BLUE)
    
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    st = doc.styles['Normal']
    st.font.name = font_name
    st.font.size = Pt(11)
    st.font.color.rgb = BLACK
    st.paragraph_format.space_after = Pt(4)

    def cp(text, bold=False, italic=False, size=11, color=None, sa=6, sb=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(sa)
        p.paragraph_format.space_before = Pt(sb)
        r = p.add_run(text)
        r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = font_name
        if color: r.font.color.rgb = color

    # Header
    title = f"ATS Analysis Report: {data.get('job_title', 'Target Role')} at {data.get('company_name', 'Target Company')}"
    cp(title, bold=True, size=16, color=primary_rgb, sa=14)
    cp("Generated based on your master resume and the job description.", italic=True, size=10, color=GRAY, sa=16)

    # Fit Score
    cp(f"ATS Fit Score: {data.get('fit_score', 0)}%", bold=True, size=14, color=secondary_rgb, sa=10)

    # Top Strengths
    cp("Top Strengths:", bold=True, size=12, color=primary_rgb, sa=4, sb=6)
    strengths = data.get('top_strengths', [])
    if not strengths:
        cp("None identified.", sa=6)
    else:
        for s in strengths:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run('\u2022 ' + str(s))
            r.font.name = font_name
            r.font.size = Pt(11)

    # Missing Keywords
    cp("Missing Keywords:", bold=True, size=12, color=RGBColor(0xC0, 0x39, 0x2B), sa=4, sb=12)
    missing = data.get('missing_keywords', [])
    if not missing:
        cp("No missing keywords identified. Great match!", sa=6)
    else:
        for m in missing:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run('\u2022 ' + str(m))
            r.font.name = font_name
            r.font.size = Pt(11)

    # Recommendations
    cp("Actionable Recommendations:", bold=True, size=12, color=primary_rgb, sa=4, sb=12)
    recommendations = data.get('recommendations', [])
    if not recommendations:
        cp("No specific recommendations generated.", sa=6)
    else:
        for rec in recommendations:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run('\u2022 ' + str(rec))
            r.font.name = font_name
            r.font.size = Pt(11)

    # Requirements Mapping Matrix
    mapping = data.get('requirements_mapping', [])
    if mapping:
        cp("Requirements Mapping Matrix:", bold=True, size=12, color=primary_rgb, sa=8, sb=16)
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'JD Requirement'
        hdr_cells[1].text = 'Resume Satisfaction'
        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.name = font_name
                    run.font.size = Pt(11)
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'EAEAEA')
            tcPr.append(shd)

        for req in mapping:
            row_cells = table.add_row().cells
            r_req = row_cells[0].paragraphs[0].add_run(str(req.get('requirement', '')))
            r_req.font.name = font_name; r_req.font.size = Pt(10)
            
            r_sat = row_cells[1].paragraphs[0].add_run(str(req.get('satisfaction', '')))
            r_sat.font.name = font_name; r_sat.font.size = Pt(10)

    filename = f"Generated_ATS_Report_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath

def generate_profile_optimization_doc(data, output_dir):
    doc = setup_doc(with_border=True, template="executive")
    font_name = 'Calibri'
    
    # Title Header
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; h.paragraph_format.space_after = Pt(2)
    r = h.add_run("CareerSuite Profile Optimization Guide")
    r.bold = True; r.font.size = Pt(22); r.font.color.rgb = NAVY; r.font.name = font_name
    
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(4)
    r = c.add_run("Tailored Platform Optimization Strategy & Copy Instructions")
    r.font.size = Pt(11); r.font.name = font_name; r.font.color.rgb = GRAY
    add_bottom_border(c)

    # 1. LINKEDIN SECTION
    heading(doc, 'LinkedIn Profile Optimization', template="executive")
    
    li = data.get('linkedin', {})
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(6)
    p_meta.paragraph_format.space_after = Pt(4)
    r = p_meta.add_run('Recommended Profile Details:\n'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    r = p_meta.add_run('• Location: '); r.bold = True; r.font.name = font_name
    r = p_meta.add_run(f"{li.get('location', 'N/A')}\n"); r.font.name = font_name
    r = p_meta.add_run('• Top Education Display: '); r.bold = True; r.font.name = font_name
    r = p_meta.add_run(f"{li.get('education_display', 'N/A')}\n"); r.font.name = font_name
    r = p_meta.add_run('• Certifications formatting: '); r.bold = True; r.font.name = font_name
    r = p_meta.add_run(f"{li.get('certifications_display', 'N/A')}\n"); r.font.name = font_name
    
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    r = p.add_run('ATS Headline Options:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    for h_opt in li.get('headline_options', []):
        p_opt = doc.add_paragraph()
        p_opt.paragraph_format.left_indent = Inches(0.25)
        p_opt.paragraph_format.space_after = Pt(2)
        r_bullet = p_opt.add_run('\u2022 ')
        r_bullet.font.name = font_name
        r_text = p_opt.add_run(h_opt)
        r_text.font.size = Pt(10); r_text.font.name = font_name; r_text.italic = True
        
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Optimized About Section:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    
    about_text = li.get('about_me', '')
    for para in about_text.split('\n'):
        if para.strip():
            p_para = doc.add_paragraph()
            p_para.paragraph_format.left_indent = Inches(0.25)
            p_para.paragraph_format.space_after = Pt(4)
            r_text = p_para.add_run(para.strip())
            r_text.font.size = Pt(10); r_text.font.name = font_name
            
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Experience Optimization:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    for exp in li.get('experience_formatting', []):
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.left_indent = Inches(0.25)
        p_exp.paragraph_format.space_before = Pt(6)
        p_exp.paragraph_format.space_after = Pt(2)
        r_title = p_exp.add_run(f"{exp.get('title', '')} at {exp.get('company', '')}")
        r_title.bold = True; r_title.font.size = Pt(11); r_title.font.name = font_name
        
        desc = exp.get('formatted_description', '')
        for d_para in desc.split('\n'):
            if d_para.strip():
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.left_indent = Inches(0.5)
                p_desc.paragraph_format.space_after = Pt(2)
                r_desc = p_desc.add_run(d_para.strip())
                r_desc.font.size = Pt(10); r_desc.font.name = font_name

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Technical Skills & Keywords to Add:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    p_skills = doc.add_paragraph()
    p_skills.paragraph_format.left_indent = Inches(0.25)
    r_skills = p_skills.add_run(", ".join(li.get('skills_recommendation', [])))
    r_skills.font.size = Pt(10); r_skills.font.name = font_name; r_skills.italic = True

    # 2. INDEED SECTION
    heading(doc, 'Indeed Resume Optimization', template="executive")
    ind = data.get('indeed', {})
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(6)
    p_meta.paragraph_format.space_after = Pt(4)
    r = p_meta.add_run('Recommended Profile Details:\n'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    r = p_meta.add_run('• Location: '); r.bold = True; r.font.name = font_name
    r = p_meta.add_run(f"{ind.get('location', 'N/A')}\n"); r.font.name = font_name
    r = p_meta.add_run('• Top Education Display: '); r.bold = True; r.font.name = font_name
    r = p_meta.add_run(f"{ind.get('education_display', 'N/A')}\n"); r.font.name = font_name
    
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    r = p.add_run('Indeed Resume Headline:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    p_opt = doc.add_paragraph()
    p_opt.paragraph_format.left_indent = Inches(0.25)
    r_text = p_opt.add_run(ind.get('headline', ''))
    r_text.font.size = Pt(10); r_text.font.name = font_name; r_text.italic = True
    
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Indeed Summary:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.left_indent = Inches(0.25)
    r_sum = p_sum.add_run(ind.get('summary', ''))
    r_sum.font.size = Pt(10); r_sum.font.name = font_name
    
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Indeed Experience descriptions:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    for exp in ind.get('experience_formatting', []):
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.left_indent = Inches(0.25)
        p_exp.paragraph_format.space_before = Pt(6)
        p_exp.paragraph_format.space_after = Pt(2)
        r_title = p_exp.add_run(f"{exp.get('title', '')} at {exp.get('company', '')}")
        r_title.bold = True; r_title.font.size = Pt(11); r_title.font.name = font_name
        
        desc = exp.get('formatted_description', '')
        for d_para in desc.split('\n'):
            if d_para.strip():
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.left_indent = Inches(0.5)
                p_desc.paragraph_format.space_after = Pt(2)
                r_desc = p_desc.add_run(d_para.strip())
                r_desc.font.size = Pt(10); r_desc.font.name = font_name

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Indeed Skills Checklist:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    p_skills = doc.add_paragraph()
    p_skills.paragraph_format.left_indent = Inches(0.25)
    r_skills = p_skills.add_run(", ".join(ind.get('skills', [])))
    r_skills.font.size = Pt(10); r_skills.font.name = font_name; r_skills.italic = True

    # 3. GUIDES SECTION
    heading(doc, 'Step-by-Step Platform Guides', template="executive")
    gd = data.get('guides', {})
    
    def add_checklist_item(target_doc, step_text):
        p_step = target_doc.add_paragraph()
        p_step.paragraph_format.left_indent = Inches(0.25)
        p_step.paragraph_format.space_after = Pt(4)
        p_step.paragraph_format.line_spacing = 1.15
        
        r_bullet = p_step.add_run('\u2610  ')
        r_bullet.font.name = font_name
        r_bullet.font.size = Pt(10)
        
        import re
        match = re.match(r'^(Step \d+:|General tip \d+:)\s*(.*)$', step_text)
        if match:
            prefix = match.group(1)
            rest = match.group(2)
            r_prefix = p_step.add_run(prefix + " ")
            r_prefix.bold = True
            r_prefix.font.name = font_name
            r_prefix.font.size = Pt(10)
            
            r_rest = p_step.add_run(rest)
            r_rest.font.name = font_name
            r_rest.font.size = Pt(10)
        else:
            r_text = p_step.add_run(step_text)
            r_text.font.name = font_name
            r_text.font.size = Pt(10)
    
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
    r = p.add_run('LinkedIn Checklist:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    for step in gd.get('linkedin', []):
        add_checklist_item(doc, step)
        
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Indeed Checklist:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    for step in gd.get('indeed', []):
        add_checklist_item(doc, step)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    r = p.add_run('Other Job Boards Advice:'); r.bold = True; r.font.size = Pt(11); r.font.name = font_name
    for step in gd.get('others', []):
        add_checklist_item(doc, step)

    filename = f"Profile_Optimization_Guide_{int(datetime.now().timestamp())}.docx"
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    return filepath
